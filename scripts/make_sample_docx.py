"""来源说明：FitPilot 知识库公开资料清单。"""

from docx import Document

doc = Document()
doc.add_heading("减脂与蛋白质安排：常见问答（公开科普整理）", level=1)
doc.add_paragraph(
    "本文档为 FitPilot 演示用 Word 语料，综合公开营养科普常见共识整理，"
    "不替代注册营养师或医生的个性化建议。"
)
doc.add_heading("减脂期蛋白质怎么吃？", level=2)
doc.add_paragraph(
    "多数健身爱好者可按体重 1.6–2.2 g/kg/天估算蛋白目标（个体差异大）。"
    "优先保证全日总量，再考虑餐次分配；力量训练日后尤其不要把蛋白砍得过低。"
)
doc.add_heading("食材选择", level=2)
doc.add_paragraph("鸡胸肉、鱼虾、蛋、奶制品、大豆制品都可轮换，避免单一极端食谱。")
doc.add_heading("需要就医的情况", level=2)
doc.add_paragraph("若存在肾病或其他医嘱限制蛋白摄入，必须遵医嘱，不宜自行提高蛋白目标。")
out = r"D:\FitPilot\knowledge_base\raw\curated\protein_fatloss_faq.docx"
doc.save(out)
print("wrote", out)
